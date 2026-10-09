"""Production contracts and recovery, with synthetic model responses (no API spend)."""
import copy
import io
import json
import tempfile
import unittest
import uuid
import wave
from pathlib import Path
from unittest.mock import patch
import production_api as production


def scene():
    return {'id':str(uuid.uuid4()),'category':'line-art','title':'Su döngüsü','duration':30,'seed':42,'speed':1,'detail':1,
            'background':'#1d2420','palette':['#f1ead8','#f2b440'],'objects':[], 'composed':True,'narration':'Su buharlaşır.'}


class ProductionTests(unittest.TestCase):
    def setUp(self):
        temporary=tempfile.TemporaryDirectory();self.addCleanup(temporary.cleanup)
        data=patch.object(production,'DATA',Path(temporary.name));data.start();self.addCleanup(data.stop)
        # Gerçek Whisper çalıştırılmaz: hizalama varsayılan olarak 'kullanılamıyor' (2. adım atlanır).
        whisper=patch.object(production.speech_sync,'align_words',side_effect=RuntimeError('whisper yok'));whisper.start();self.addCleanup(whisper.stop)

    def job(self):
        identifier=str(uuid.uuid4());production.write(production.folder(identifier)/'job.json',{'id':identifier,'state':'running'})
        return identifier

    def plan(self, count=1):
        return {'title':'Ders','theme':'chalk','scenes':[{'title':f'Sahne {i}','category':'line-art','duration':30,
            'narration':'Su buharlaşır ve atmosfere yükselir.','visual':'Deniz ve güneş çiz, oklarla buharlaşmayı göster.','sourceRefs':[]} for i in range(count)]}

    def test_document_and_word_order(self):
        from docx import Document
        doc=Document();doc.add_paragraph('Birinci bölüm: suyun buharlaşması anlatılır.');doc.add_paragraph('İkinci bölüm: yoğuşma ile bulutlar oluşur.')
        stream=io.BytesIO();doc.save(stream)
        parsed=production.document('ders.docx',stream.getvalue())
        self.assertLess(parsed['text'].index('Birinci'),parsed['text'].index('İkinci'))
        self.assertEqual(len(parsed['sections']),2)
        self.assertEqual(len(parsed['sha256']),64)

    def test_empty_and_unsupported_documents(self):
        for name,content in [('x.txt',b''),('x.exe',b'not a document'),('x.txt',b'Hi')]:
            with self.subTest(name=name),self.assertRaises(ValueError):production.document(name,content)

    def test_rejects_invented_source_refs(self):
        doc=production.document('source.txt',b'A sufficient educational source document for this test.')
        plan=self.plan();plan['source']={'id':doc['id']};plan['scenes'][0]['sourceRefs']=['sayfa 999']
        with self.assertRaises(ValueError):production.validate_plan(plan)
        plan['scenes'][0]['sourceRefs']=['metin'];self.assertEqual(production.validate_plan(plan)['source']['name'],'source.txt')

    def test_long_project_and_limits(self):
        project={'version':1,'name':'Uzun ders','scenes':[scene() for _ in range(360)]}
        clean=production.validate_project(project)
        self.assertEqual(sum(s['duration'] for s in clean['scenes']),10800)
        project['scenes'].append(scene())
        with self.assertRaises(ValueError):production.validate_project(project)
        project['scenes']=[scene()];project['scenes'][0]['duration']=float('nan')
        with self.assertRaises(ValueError):production.validate_project(project)

    def test_normalizes_plan_to_target_duration(self):
        identifier=self.job();plan=self.plan(7)
        with patch.object(production.ai,'call_json',return_value=(plan,{'model':'glm-5.3'})):
            result=production.job_plan(identifier,{'durationSeconds':181,'prompt':'Su döngüsünü öğret','theme':'chalk','category':'line-art'})
        self.assertAlmostEqual(sum(s['duration'] for s in result['plan']['scenes']),181,places=3)

    def test_completed_scenes_survive_failure_and_resume(self):
        plan=self.plan(3);identifier=self.job()
        with patch.object(production.ai,'generate',side_effect=[(200,{'scene':scene()}),(502,{'error':'quota'})]):
            with self.assertRaises(production.ai.ModelError):production.job_build(identifier,{'plan':plan})
        partial=production.read(production.folder(identifier)/'project.json')
        self.assertEqual(len(partial['scenes']),1)
        next_job=self.job()
        with patch.object(production.ai,'generate',side_effect=[(200,{'scene':scene()}),(200,{'scene':scene()})]) as call:
            result=production.job_build(next_job,{'plan':plan,'resumeJobId':identifier})
            self.assertEqual(call.call_count,2)
        self.assertEqual(len(result['project']['scenes']),3)
        self.assertEqual(result['project']['scenes'][0]['id'],partial['scenes'][0]['id'])

    def test_cancel_and_path_traversal(self):
        identifier=self.job();production.update(identifier,cancelRequested=True)
        with self.assertRaises(InterruptedError):production.check_cancel(identifier)
        for value in ('../.env','not-a-uuid'):
            with self.assertRaises(ValueError):production.folder(value)
        with self.assertRaises(ValueError):production.asset_path('/api/animation/assets/'+identifier+'/../.env')

    def test_upload_audio_checks_magic_bytes(self):
        with self.assertRaises(ValueError):production.upload_audio('a.wav',b'not a wav file at all')
        with self.assertRaises(ValueError):production.upload_audio('a.exe',b'RIFF\x00\x00\x00\x00WAVE')

    def test_every_scene_retains_generated_audio(self):
        identifier=self.job();project={'version':1,'scenes':[scene() for _ in range(3)]}
        def synthesize(job,args,**kwargs):
            payload=production.read(Path(args[-1]))
            with wave.open(payload['output'],'wb') as wav:
                wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(8000);wav.writeframes(b'\x00\x00'*16000)
        with patch.object(production,'status',return_value={'voices':[{'id':'windows','available':True}]}),patch.object(production,'run_process',side_effect=synthesize):
            result=production.job_voice(identifier,{'project':project,'provider':'windows'})
        self.assertEqual(len(result['project']['scenes']),3)
        self.assertTrue(all('audio' in s for s in result['project']['scenes']))
        self.assertEqual(len({s['audio']['url'] for s in result['project']['scenes']}),3)

    def fake_cartesia(self,seconds):
        """cartesia_tts.py yerine: hedef WAV'ı (son argüman) yazar; ağ/API çağrısı yok."""
        calls=[]
        def synthesize(job,args,**kwargs):
            self.assertTrue(args[3].endswith('cartesia_tts.py'));calls.append(args[-2])
            with wave.open(args[-1],'wb') as wav:
                wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(8000);wav.writeframes(b'\x00\x00'*int(8000*seconds))
        return calls,synthesize

    def test_single_scene_voice_only_touches_that_scene(self):
        identifier=self.job();scenes=[scene() for _ in range(3)]
        for i,s in enumerate(scenes):s.update(duration=5,narration=f'Sahne {i} anlatımı.')
        scenes[1]['objects']=[{'type':'circle','x':.5,'y':.5,'width':.1,'height':.1,'color':'#f2b440','lineWidth':1,'start':4,'duration':1,'text':'','points':[],'motion':'fade'}]
        project=production.validate_project({'version':1,'scenes':scenes})
        calls,synthesize=self.fake_cartesia(12)
        with patch.object(production,'status',return_value={'voices':[{'id':'windows','available':False},{'id':'cartesia','available':True}]}),patch.object(production,'run_process',side_effect=synthesize):
            result=production.job_voice(identifier,{'project':project,'provider':'cartesia','sceneIndex':1,'sceneId':project['scenes'][1]['id']})
        self.assertEqual(calls,['Sahne 1 anlatımı.'])
        out=result['project']['scenes']
        self.assertNotIn('audio',out[0]);self.assertNotIn('audio',out[2])
        self.assertEqual([out[0]['duration'],out[2]['duration']],[5,5])
        self.assertEqual(out[0]['narration'],'Sahne 0 anlatımı.')
        self.assertEqual(out[1]['audio']['provider'],'cartesia');self.assertAlmostEqual(out[1]['audio']['duration'],12,places=2)
        self.assertAlmostEqual(out[1]['duration'],12.6,places=3)                     # süre sese uyar (ses + 0,6)
        self.assertAlmostEqual(out[1]['objects'][0]['start'],round(4*12.6/5,3),places=3)  # nesne zamanları ölçeklenir
        self.assertFalse(result['aligned']);self.assertEqual(out[1]['words'],[])        # Whisper yoksa hizalama atlanır
        self.assertEqual((result['sceneIndex'],result['sceneId'],result['scene']['id']),(1,out[1]['id'],out[1]['id']))
        # sceneId tek başına da seçer
        calls,synthesize=self.fake_cartesia(2)
        with patch.object(production,'status',return_value={'voices':[{'id':'cartesia','available':True}]}),patch.object(production,'run_process',side_effect=synthesize):
            result=production.job_voice(self.job(),{'project':project,'provider':'cartesia','sceneId':project['scenes'][2]['id']})
        self.assertEqual(calls,['Sahne 2 anlatımı.']);self.assertEqual(result['sceneIndex'],2)
        self.assertEqual(result['project']['scenes'][2]['duration'],2.6)             # kısa ses süreyi de kısaltır
        self.assertEqual(result['project']['scenes'][1]['duration'],5)               # diğer sahneler değişmez
        calls,synthesize=self.fake_cartesia(2)
        with patch.object(production,'status',return_value={'voices':[{'id':'cartesia','available':True}]}),patch.object(production,'run_process',side_effect=synthesize):
            result=production.job_voice(self.job(),{'project':project,'provider':'cartesia','sceneIndex':2,'fitToAudio':False})
        self.assertEqual(result['project']['scenes'][2]['duration'],5)               # fitToAudio=False: yalnızca uzatır

    def test_single_scene_voice_rejects_bad_selection(self):
        project={'version':1,'scenes':[scene() for _ in range(2)]};project['scenes'][1]['narration']=' '
        bad=[{'sceneIndex':-1},{'sceneIndex':2},{'sceneIndex':True},{'sceneIndex':'0'},{'sceneIndex':0.0},
             {'sceneId':'yok'},{'sceneId':7},{'sceneIndex':0,'sceneId':project['scenes'][1]['id']},{'sceneIndex':1}]
        with patch.object(production,'status',return_value={'voices':[{'id':'cartesia','available':True}]}),patch.object(production.POOL,'submit') as submit:
            for extra in bad:
                with self.subTest(extra=extra),self.assertRaises(ValueError):
                    production.create_job({'type':'voice','project':project,'provider':'cartesia',**extra})
            with self.assertRaises(ValueError):production.create_job({'type':'voice','project':project,'provider':'nope','sceneIndex':0})
            submit.assert_not_called()
            production.create_job({'type':'voice','project':project,'provider':'cartesia','sceneIndex':0})
            self.assertEqual(submit.call_count,1)


if __name__=='__main__':unittest.main()
